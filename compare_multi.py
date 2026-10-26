# -*- coding: utf-8 -*-
"""Reusable Yakjin Aerie order-vs-PO comparison. Requires: pip install openpyxl"""
import argparse, re
from collections import defaultdict, Counter
from datetime import datetime, timedelta
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def excel_rows(path, preferred, cols):
    wb=load_workbook(path, read_only=True, data_only=True)
    sh=wb[preferred] if preferred in wb.sheetnames else wb.worksheets[0]
    # Some PO exports incorrectly declare worksheet dimensions as A1:A1.
    # Ignore that metadata so openpyxl reads the actual XML rows.
    if hasattr(sh, 'reset_dimensions'): sh.reset_dimensions()
    rows=[]
    for row in sh.iter_rows(values_only=True):
        values=list(row[:cols])
        rows.append(values+[None]*(cols-len(values)))
    wb.close()
    return rows

def t(x):return str(x or '').strip()
def po7(x):return bool(re.fullmatch(r'[0-9]{7}',t(x)))
def date(x):
    if isinstance(x,(int,float)):
        return (datetime(1899,12,30)+timedelta(days=float(x))).strftime('%Y-%m-%d')
    if isinstance(x,datetime):return x.strftime('%Y-%m-%d')
    s=t(x)
    for fmt in ('%m/%d/%y','%m/%d/%Y','%Y-%m-%d','%d/%m/%Y'):
        try:return datetime.strptime(s,fmt).strftime('%Y-%m-%d')
        except ValueError:pass
    return s

def season(x):
    s=t(x).upper()
    year=re.search(r'20(\d\d)',s)
    yy=year.group(1) if year else (re.search(r'\b(\d\d)\b',s).group(1) if re.search(r'\b(\d\d)\b',s) else '')
    if 'SUM' in s or re.search(r'\bSU\b',s):return 'SU'+yy
    if 'SPR' in s or re.search(r'\bSP\b',s):return 'SP'+yy
    if 'HOL' in s:return 'HOL'+yy
    if 'BTS' in s:return 'BTS'+yy
    return s

def factory(x):
    s=re.sub(r'[^A-Z0-9]+',' ',t(x).upper()).strip()
    if 'VINA KOREA' in s:return 'VINA KOREA'
    if 'YAKJIN JAYA' in s:return 'PT YAKJIN JAYA'
    if 'YAKJIN VIETNAM' in s:return 'YAKJIN VIETNAM'
    if 'YAKJIN CAMBO' in s:return 'YAKJIN CAMBODIA'
    return s

def quantity(x):
    try:return int(round(float(x or 0)))
    except (TypeError,ValueError):return 0

def output_sheet(wb,name,headers,records,widths):
    sh=wb.create_sheet(name)
    sh.append(headers)
    for row in records: sh.append([str(v) if isinstance(v,(list,tuple,set)) else v for v in row])
    for cell in sh[1]:
        cell.fill=PatternFill('solid',fgColor='17365D')
        cell.font=Font(name='Aptos Narrow',size=10,bold=True,color='FFFFFF')
        cell.alignment=Alignment(vertical='center')
    sh.row_dimensions[1].height=29
    for i,w in enumerate(widths,1): sh.column_dimensions[get_column_letter(i)].width=max(12,w/1.3)
    sh.freeze_panes='A2'
    if records:
        table=Table(displayName=re.sub(r'\W','',name)+'Table',ref=f'A1:{get_column_letter(len(headers))}{len(records)+1}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        sh.add_table(table)
    return sh


def parse_order(path):
    rows=excel_rows(path,'Sheet1',51)
    valid=[];excluded=[]
    for r in rows[2:]:
        if not t(r[24]):continue
        item=[season(r[18]),t(r[7]),t(r[15]),t(r[17]),t(r[24]),t(r[42]).zfill(3),t(r[40]),date(r[30]),factory(r[3]),quantity(r[44])]
        if po7(r[24]):valid.append(item)
        else:excluded.append(item+['PO# not exactly 7 digits'])
    return valid,excluded

def parse_po(path):
    rows=excel_rows(path,'Sheet0',33)
    out=[]
    for r in rows[1:]:
        if not t(r[14]) or not t(r[17]):continue
        out.append([season(r[10]),t(r[6]),t(r[17]),t(r[18]),t(r[14]),t(r[19]).zfill(3),date(r[23]),factory(r[29]),quantity(r[21])])
    if not out:
        raise ValueError(f'PO Detailed View 파일에서 유효한 PO 데이터를 찾지 못했습니다: {path.name}. Sheet0에 상세 데이터가 있는 원본 파일을 선택하세요 (피벗 요약/빈 파일 제외).')
    return out

def run_multi(order_paths,po_paths,out_path):
    # File sequence is oldest to newest; latest upload overrides the entire
    # PO/style/color group (preserving multiple channels and dates within a file).
    # PO keys are season-scoped to avoid accidental cross-season overwrite.
    def combine(paths,parser):
        latest={}
        audits=[]
        for idx,path in enumerate(paths):
            parsed=parser(path)
            rows=parsed[0] if isinstance(parsed,tuple) else parsed
            groups=defaultdict(list)
            for r in rows:groups[(r[0],r[2],r[4],r[5])].append(r)
            replaced=0
            for k,rr in groups.items():
                if k in latest:replaced+=1
                latest[k]=(idx,path.name,rr)
            audits.append([path.name,idx+1,len(rows),len(groups),replaced])
        merged=[r for _,_,rr in latest.values() for r in rr]
        return merged,audits,latest
    order,oa,ol=combine(order_paths,parse_order)
    if not order: raise ValueError('오더 수주 현황에서 유효한 7자리 PO 데이터가 없습니다. 파일 형식을 확인하세요.')
    po,pa,pl=combine(po_paths,parse_po)
    excluded=[]
    for idx,path in enumerate(order_paths):
        excluded.extend([r+[path.name] for r in parse_order(path)[1]])
    # Same style/PO/color is the identity for reconciliation.
    def identity(r):return (r[2],r[4],r[5])
    om=defaultdict(list);pm=defaultdict(list)
    color_names=defaultdict(set);bycode=defaultdict(set)
    style_names=defaultdict(set);po_style_names=defaultdict(set)
    for r in order:
        om[identity(r)].append(r)
        if r[6]:color_names[(r[2],r[5])].add(r[6]);bycode[r[5]].add(r[6])
        if r[3]:style_names[r[2]].add(r[3])
    for r in po:
        pm[identity(r)].append(r)
        if r[3]:po_style_names[r[2]].add(r[3])
    detail=[];review=[];stat=Counter()
    summaries={k:defaultdict(lambda:[0,0,0]) for k in ('Factory','Season','Division','Style','Month')}
    for k in sorted(set(om)|set(pm)):
        a=om.get(k,[]);b=pm.get(k,[])
        oq=sum(x[9] for x in a);pq=sum(x[8] for x in b)
        da=sorted(set(x[7] for x in a));db=sorted(set(x[6] for x in b))
        fa=sorted(set(x[8] for x in a));fb=sorted(set(x[7] for x in b))
        sa=sorted(set(x[0] for x in a));sb=sorted(set(x[0] for x in b))
        div=sorted(set(x[1] for x in a)) or sorted(set(x[1] for x in b))
        sn=style_names[k[0]] or po_style_names[k[0]]
        description=next(iter(sn)) if len(sn)==1 else (' / '.join(sorted(sn)) if sn else '')
        names=color_names[(k[0],k[2])] or bycode[k[2]]
        cname=next(iter(names)) if len(names)==1 else ('COLOR NAME REVIEW REQUIRED' if len(names)>1 else '')
        if len(names)>1:review.append([k[0],k[2],', '.join(sorted(names))])
        if not a:status='ONLY PO'
        elif not b:status='ONLY ORDER'
        elif oq!=pq:status='QTY DIFF'
        elif da!=db:status='DATE CHANGE'
        elif fa!=fb:status='FACTORY CHANGE'
        elif sa!=sb:status='SEASON CHANGE'
        else:status='MATCH'
        stat[status]+=1
        detail.append([', '.join(sa or sb),', '.join(div),k[0],description,k[1],k[2],cname,', '.join(da),', '.join(db),', '.join(fa),', '.join(fb),oq,pq,oq-pq,status,', '.join(sb)])
        for r in a:
            for cat,val in [('Factory',r[8]),('Season',r[0]),('Division',r[1]),('Style',r[2]+' | '+description),('Month',r[7][:7])]:
                summaries[cat][val][0]+=r[9]
        for r in b:
            for cat,val in [('Factory',r[7]),('Season',r[0]),('Division',r[1]),('Style',r[2]+' | '+description),('Month',r[6][:7])]:
                summaries[cat][val][1]+=r[8]
        if status!='MATCH':
            cats={'Factory':set(fa+fb),'Season':set(sa+sb),'Division':set(div),'Style':{k[0]+' | '+description},'Month':set(x[:7] for x in da+db)}
            for cat,vals in cats.items():
                for val in vals:summaries[cat][val][2]+=1
    wb=Workbook();wb.remove(wb.active)
    total_o=sum(x[9] for x in order);total_p=sum(x[8] for x in po)
    output_sheet(wb,'Dashboard',['Metric','Value'],[['Order Qty (7-digit PO)',total_o],['PO Detailed Qty',total_p],['Net Difference',total_o-total_p],['Order Files',len(order_paths)],['PO Files',len(po_paths)],['Compared Keys',len(detail)],['Excluded Order Rows',len(excluded)]]+[[s,stat[s]] for s in ['MATCH','QTY DIFF','DATE CHANGE','FACTORY CHANGE','SEASON CHANGE','ONLY ORDER','ONLY PO']],[35,21])
    output_sheet(wb,'Discrepancy',['Season','Division','Style#','Style Description','PO#','Color Code','Color Name (Order)','Ex-Fty','H/O','Order Factory','PO Factory','Order Qty','PO Qty','Difference','Status','PO Season'],detail,[18,31,15,38,16,16,34,21,21,27,27,16,16,16,22,18])
    for cat,values in summaries.items():
        output_sheet(wb,cat+' Summary',[cat,'Order Qty','PO Qty','Difference','Non-match keys'],[[key,v[0],v[1],v[0]-v[1],v[2]] for key,v in sorted(values.items())],[42,19,19,19,21])
    output_sheet(wb,'Upload Audit',['File','Upload Sequence','Rows Read','Unique Keys','Replaced Keys'],[['ORDER: '+x[0]]+x[1:] for x in oa]+[['PO: '+x[0]]+x[1:] for x in pa],[53,21,17,18,19])
    output_sheet(wb,'Color Review',['Style#','Color Code','Conflicting Names'],review,[16,17,80])
    output_sheet(wb,'Excluded Orders',['Season','Division','Style#','Style Description','PO#','Color Code','Color Name','Ex-Fty','Factory','Qty','Reason','Source File'],excluded,[16,30,15,36,17,16,30,19,27,15,28,45])
    output_sheet(wb,'Order Source',['Season','Division','Style#','Style Description','PO#','Color Code','Color Name','Ex-Fty','Factory','Order Qty'],order,[16,30,15,36,17,16,30,19,27,17])
    output_sheet(wb,'PO Source',['Season','Division','Style#','Style Description','PO#','Color Code','H/O','Factory','PO Qty'],po,[16,30,15,36,17,16,19,27,17])
    wb.save(str(out_path))
    return {'order_qty':total_o,'po_qty':total_p,'keys':len(detail),'statuses':dict(stat)}
